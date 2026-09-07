# Geometry QA for the deck: flag likely text overflow, out-of-slide shapes, text-text overlaps.
# Heuristic char widths: Segoe UI ~0.52*pt, Consolas ~0.55*pt (px-free, all in points).
import math, sys
from pptx import Presentation
from pptx.util import Emu

PATH = r"E:\agent\esp32_lua_llm\PowerPoint sharing\ESP32_Lua_LLM_Architecture.pptx"
SLIDE_W, SLIDE_H = 13.333, 7.5
def inches(v): return Emu(v).inches if v is not None else None

def para_font(para, default=13.0, default_face="Segoe UI"):
    size, face = None, None
    for r in para.runs:
        if r.font.size: size = r.font.size.pt
        if r.font.name: face = r.font.name
        if size: break
    return (size or default), (face or default_face)

def est_text_h(shape):
    tf = shape.text_frame
    w_in = inches(shape.width)
    total_h = 0.0
    max_line_w = 0.0
    for para in tf.paragraphs:
        text = "".join(r.text for r in para.runs)
        if not text.strip():
            total_h += 0.12
            continue
        fs, face = para_font(para)
        cw = 0.55 * fs / 72.0 if face == "Consolas" else 0.52 * fs / 72.0  # char width in inches
        chars_per_line = max(1, int((w_in - 0.1) / cw))
        # account for explicit line breaks in text
        lines = 0
        for seg in text.split("\n"):
            lines += max(1, math.ceil(len(seg) / chars_per_line))
            max_line_w = max(max_line_w, min(len(seg), chars_per_line) * cw)
        total_h += lines * fs * 1.22 / 72.0
    return total_h, max_line_w

prs = Presentation(PATH)
issues = []
for si, slide in enumerate(prs.slides, 1):
    boxes = []
    for sp in slide.shapes:
        x, y = inches(sp.left), inches(sp.top)
        w, h = inches(sp.width), inches(sp.height)
        if x is None: continue
        if x < -0.05 or y < -0.05 or x + w > SLIDE_W + 0.05 or y + h > SLIDE_H + 0.05:
            issues.append(f"S{si}: OUT-OF-SLIDE {sp.shape_type} at ({x:.2f},{y:.2f},{w:.2f},{h:.2f}) text={getattr(sp,'text','')[:40]!r}")
        if sp.has_text_frame and sp.text.strip():
            eh, mw = est_text_h(sp)
            if eh > h * 1.12 + 0.06:
                issues.append(f"S{si}: OVERFLOW? est {eh:.2f}in > box {h:.2f}in at ({x:.2f},{y:.2f}) text={sp.text[:60]!r}")
            boxes.append((x, y, w, h, sp.text[:30]))
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            a, b = boxes[i], boxes[j]
            ix = max(0, min(a[0]+a[2], b[0]+b[2]) - max(a[0], b[0]))
            iy = max(0, min(a[1]+a[3], b[1]+b[3]) - max(a[1], b[1]))
            inter = ix * iy
            amin = min(a[2]*a[3], b[2]*b[3])
            if inter > 0.25 * amin and inter > 0.15:
                issues.append(f"S{si}: OVERLAP {a[4]!r} <-> {b[4]!r} ({inter:.2f} sq in)")

print(f"{len(prs.slides)} slides checked")
if issues:
    print("\n".join(issues))
else:
    print("no geometry issues found")
